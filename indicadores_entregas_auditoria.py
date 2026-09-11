from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st
from supabase_client import get_client

INDICADOR = "ENTREGAS NO PRAZO"
TZ_APP = ZoneInfo("America/Sao_Paulo")


def _agora_local():
    return datetime.now(TZ_APP)


def _norm_text(value):
    return "" if value is None else str(value).strip()


def _norm_code(value):
    s = _norm_text(value)
    if not s or s.lower() == "nan":
        return ""
    if s.endswith(".0"):
        s = s[:-2]
    return s.upper()


def _find_column(df, *names):
    def key(v):
        import unicodedata
        s = unicodedata.normalize("NFKD", str(v))
        s = "".join(ch for ch in s if not unicodedata.combining(ch))
        return " ".join(s.strip().upper().split())
    lookup = {key(c): c for c in df.columns}
    for name in names:
        if key(name) in lookup:
            return lookup[key(name)]
    raise ValueError(f"Coluna não encontrada: {names[0]}")


def _sha256_bytes(data: bytes):
    return hashlib.sha256(data).hexdigest()


def _read_uploaded_excel(uploaded, **kwargs):
    return pd.read_excel(BytesIO(uploaded.getvalue()), engine="openpyxl", **kwargs)


def _format_date_series(series):
    return series.dt.strftime("%d/%m/%Y").where(series.notna(), "")


def calcular_entregas_no_prazo_auditoria(relatorio_file, cadastros_file, data_analise):
    rel_bytes = relatorio_file.getvalue()
    cad_bytes = cadastros_file.getvalue()

    rel = _read_uploaded_excel(relatorio_file, sheet_name="RelatorioTratado", dtype=str)
    cad = _read_uploaded_excel(cadastros_file, sheet_name=0, header=1, dtype=str)

    c_projeto = _find_column(rel, "Projeto")
    c_codigo = _find_column(rel, "Código", "Codigo")
    c_ultima = _find_column(rel, "Última solicitação", "Ultima solicitacao")
    c_pend = _find_column(rel, "Pendência", "Pendencia")
    c_sep = _find_column(rel, "Data de separação", "Data de separacao")
    c_cm = _find_column(rel, "DATA CM")
    c_semana = _find_column(rel, "SEMANA DE NECESSIDADE")

    cad_codigo = _find_column(cad, "Código", "Codigo")
    cad_tipo = _find_column(cad, "Tipo")

    cad_base = cad[[cad_codigo, cad_tipo]].copy()
    cad_base["codigo_norm"] = cad_base[cad_codigo].map(_norm_code)
    cad_base["tipo"] = cad_base[cad_tipo].fillna("").astype(str).str.strip().str.upper()
    cad_base = cad_base[cad_base["codigo_norm"] != ""].drop_duplicates("codigo_norm", keep="last")
    tipos = cad_base.set_index("codigo_norm")["tipo"].to_dict()

    base = rel[[c_projeto, c_codigo, c_ultima, c_pend, c_sep, c_cm, c_semana]].copy()
    base.columns = ["projeto", "codigo", "ultima_solicitacao", "pendencia", "data_separacao", "data_cm", "semana_necessidade"]
    base["linha_origem"] = base.index + 2
    base["projeto"] = base["projeto"].fillna("").astype(str).str.strip()
    base["codigo_norm"] = base["codigo"].map(_norm_code)
    base["tipo"] = base["codigo_norm"].map(tipos).fillna("")
    base["ultima_solicitacao"] = pd.to_datetime(base["ultima_solicitacao"], dayfirst=True, errors="coerce")
    base["data_separacao"] = pd.to_datetime(base["data_separacao"], dayfirst=True, errors="coerce")
    base["data_cm"] = pd.to_datetime(base["data_cm"], dayfirst=True, errors="coerce")
    base["pendencia"] = pd.to_numeric(
        base["pendencia"].fillna("0").astype(str).str.replace(",", ".", regex=False), errors="coerce"
    ).fillna(0.0)
    base["semana_num"] = pd.to_numeric(
        base["semana_necessidade"].fillna("").astype(str).str.replace(",", ".", regex=False), errors="coerce"
    )

    inicio = pd.Timestamp(data_analise.replace(day=1))
    fim = pd.Timestamp(data_analise)
    periodo = base[base["data_cm"].between(inicio, fim, inclusive="both")].copy()
    if periodo.empty:
        raise ValueError(
            f"Nenhuma linha com DATA CM entre {inicio.strftime('%d/%m/%Y')} e {fim.strftime('%d/%m/%Y')}."
        )

    periodo["pendencia_bruta"] = periodo["pendencia"] > 0
    periodo["semana_valida"] = periodo["semana_num"].notna()
    periodo["tipo_ii"] = periodo["tipo"].eq("II")
    periodo["zerada_sem_semana"] = periodo["pendencia_bruta"] & ~periodo["semana_valida"]
    periodo["zerada_tipo_ii"] = periodo["pendencia_bruta"] & periodo["semana_valida"] & periodo["tipo_ii"]
    periodo["pendencia_efetiva"] = periodo["pendencia_bruta"] & periodo["semana_valida"] & ~periodo["tipo_ii"]
    periodo["pendencia_final"] = periodo["pendencia"].where(periodo["pendencia_efetiva"], 0.0)

    sep_preenchida = periodo["data_separacao"].notna()
    periodo["entrega_direta"] = sep_preenchida & (periodo["data_separacao"] <= periodo["data_cm"])
    periodo["excecao_ultima_solicitacao"] = (
        sep_preenchida
        & (periodo["data_separacao"] > periodo["data_cm"])
        & periodo["ultima_solicitacao"].notna()
        & (periodo["ultima_solicitacao"] > periodo["data_cm"])
    )
    periodo["sem_separacao_tipo_ii"] = periodo["data_separacao"].isna() & periodo["tipo_ii"]
    periodo["entregue_em_dia"] = (
        periodo["entrega_direta"] | periodo["excecao_ultima_solicitacao"] | periodo["sem_separacao_tipo_ii"]
    )

    def motivo_entrega(row):
        if row["entrega_direta"]:
            return "DATA SEPARAÇÃO <= DATA CM"
        if row["excecao_ultima_solicitacao"]:
            return "EXCEÇÃO: ÚLTIMA SOLICITAÇÃO > DATA CM"
        if row["sem_separacao_tipo_ii"]:
            return "SEM DATA DE SEPARAÇÃO + TIPO II"
        if pd.isna(row["data_separacao"]):
            return "NÃO ENTREGUE: SEM DATA DE SEPARAÇÃO"
        return "FORA DO PRAZO"

    def motivo_pendencia(row):
        if not row["pendencia_bruta"]:
            return "SEM PENDÊNCIA"
        if row["zerada_sem_semana"]:
            return "ZERADA: SEM SEMANA DE NECESSIDADE NUMÉRICA"
        if row["zerada_tipo_ii"]:
            return "ZERADA: MATERIAL TIPO II"
        return "PENDÊNCIA EFETIVA"

    periodo["motivo_entrega"] = periodo.apply(motivo_entrega, axis=1)
    periodo["motivo_pendencia"] = periodo.apply(motivo_pendencia, axis=1)
    periodo["codigo_encontrado_cadastro"] = (periodo["codigo_norm"] != "") & (periodo["tipo"] != "")

    projetos_validos = periodo[periodo["projeto"] != ""].copy()
    programados = int(projetos_validos["projeto"].nunique())
    solicitacoes = int(len(periodo))

    resumo_projetos = (
        projetos_validos.groupby("projeto", as_index=False)
        .agg(
            solicitacoes=("codigo", "size"),
            pendencias_brutas=("pendencia_bruta", "sum"),
            zeradas_sem_semana=("zerada_sem_semana", "sum"),
            zeradas_tipo_ii=("zerada_tipo_ii", "sum"),
            pendencias_efetivas=("pendencia_efetiva", "sum"),
            entregas_em_dia=("entregue_em_dia", "sum"),
        )
    )
    resumo_projetos["status"] = resumo_projetos["pendencias_efetivas"].map(
        lambda x: "ATENDIDO 100%" if int(x) == 0 else "PENDENTE"
    )
    resumo_projetos["pct_entregas_em_dia"] = (
        resumo_projetos["entregas_em_dia"] / resumo_projetos["solicitacoes"] * 100.0
    ).round(2)
    atendidos_100 = int((resumo_projetos["pendencias_efetivas"] == 0).sum())
    entregues_em_dia = int(periodo["entregue_em_dia"].sum())

    projetos_atendidos_pct = (atendidos_100 / programados * 100.0) if programados else 0.0
    entregas_em_dia_pct = (entregues_em_dia / solicitacoes * 100.0) if solicitacoes else 0.0
    resultado_final = projetos_atendidos_pct * entregas_em_dia_pct / 100.0

    auditoria = {
        "periodo_inicio": inicio.date().isoformat(),
        "periodo_fim": fim.date().isoformat(),
        "data_analise": data_analise.isoformat(),
        "gerado_em": _agora_local().isoformat(),
        "programados": programados,
        "atendidos_100": atendidos_100,
        "solicitacoes": solicitacoes,
        "entregues_em_dia": entregues_em_dia,
        "projetos_atendidos_pct": projetos_atendidos_pct,
        "entregas_em_dia_pct": entregas_em_dia_pct,
        "resultado_final": resultado_final,
        "entrega_direta": int(periodo["entrega_direta"].sum()),
        "excecao_ultima_solicitacao": int(periodo["excecao_ultima_solicitacao"].sum()),
        "sem_separacao_tipo_ii": int(periodo["sem_separacao_tipo_ii"].sum()),
        "pendencias_brutas": int(periodo["pendencia_bruta"].sum()),
        "pendencias_sem_semana": int(periodo["zerada_sem_semana"].sum()),
        "pendencias_tipo_ii": int(periodo["zerada_tipo_ii"].sum()),
        "pendencias_efetivas": int(periodo["pendencia_efetiva"].sum()),
        "codigos_sem_cadastro": int(((periodo["codigo_norm"] != "") & (periodo["tipo"] == "")).sum()),
        "relatorio_nome": getattr(relatorio_file, "name", "RelatorioGeral_Tratado.xlsx"),
        "relatorio_sha256": _sha256_bytes(rel_bytes),
        "relatorio_tamanho_bytes": len(rel_bytes),
        "cadastros_nome": getattr(cadastros_file, "name", "CADASTROS.xltx"),
        "cadastros_sha256": _sha256_bytes(cad_bytes),
        "cadastros_tamanho_bytes": len(cad_bytes),
    }

    detalhe = periodo.copy()
    detalhe["ultima_solicitacao"] = _format_date_series(detalhe["ultima_solicitacao"])
    detalhe["data_separacao"] = _format_date_series(detalhe["data_separacao"])
    detalhe["data_cm"] = _format_date_series(detalhe["data_cm"])
    detalhe["entregue_em_dia"] = detalhe["entregue_em_dia"].map({True: "SIM", False: "NÃO"})
    detalhe["pendencia_efetiva"] = detalhe["pendencia_efetiva"].map({True: "SIM", False: "NÃO"})
    detalhe["codigo_encontrado_cadastro"] = detalhe["codigo_encontrado_cadastro"].map({True: "SIM", False: "NÃO"})

    detalhe = detalhe[[
        "linha_origem", "projeto", "codigo", "codigo_norm", "tipo", "codigo_encontrado_cadastro",
        "ultima_solicitacao", "pendencia", "semana_necessidade", "data_separacao", "data_cm",
        "motivo_pendencia", "pendencia_final", "pendencia_efetiva", "motivo_entrega", "entregue_em_dia",
    ]].copy()
    detalhe.columns = [
        "Linha origem", "Projeto", "Código", "Código normalizado", "Tipo", "Código localizado em CADASTROS",
        "Última solicitação", "Pendência original", "Semana de necessidade", "Data de separação", "Data CM",
        "Tratamento da pendência", "Pendência final", "Pendência efetiva", "Critério de entrega", "Entregue em dia",
    ]

    excecoes = detalhe[
        detalhe["Tratamento da pendência"].str.startswith("ZERADA", na=False)
        | detalhe["Critério de entrega"].str.startswith("EXCEÇÃO", na=False)
        | detalhe["Critério de entrega"].str.contains("TIPO II", na=False)
        | detalhe["Código localizado em CADASTROS"].eq("NÃO")
    ].copy()

    cadastro_usado = (
        periodo[["codigo_norm", "tipo"]]
        .drop_duplicates()
        .sort_values(["codigo_norm", "tipo"])
        .rename(columns={"codigo_norm": "Código", "tipo": "Tipo"})
    )

    return {
        **auditoria,
        "projetos": resumo_projetos.sort_values(["status", "projeto"]).to_dict("records"),
        "detalhamento": detalhe.to_dict("records"),
        "excecoes": excecoes.to_dict("records"),
        "cadastro_usado": cadastro_usado.to_dict("records"),
    }


def _ultima_meta(indicadores):
    candidatos = [
        row for row in (indicadores or [])
        if str(row.get("indicador") or "").strip().upper() == INDICADOR and row.get("meta") is not None
    ]
    if not candidatos:
        return 0.0
    candidatos.sort(key=lambda r: pd.to_datetime(r.get("competencia"), errors="coerce"))
    try:
        return float(candidatos[-1].get("meta"))
    except Exception:
        return 0.0


def _gerar_excel_auditoria(resultado, meta):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    output = BytesIO()
    resumo_rows = [
        ["RELATÓRIO DE AUDITORIA — ENTREGAS NO PRAZO", ""],
        ["Indicador", INDICADOR],
        ["Data da análise", pd.to_datetime(resultado["data_analise"]).strftime("%d/%m/%Y")],
        ["Período inicial", pd.to_datetime(resultado["periodo_inicio"]).strftime("%d/%m/%Y")],
        ["Período final", pd.to_datetime(resultado["periodo_fim"]).strftime("%d/%m/%Y")],
        ["Gerado em", pd.to_datetime(resultado["gerado_em"]).strftime("%d/%m/%Y %H:%M:%S")],
        ["", ""],
        ["Programados", resultado["programados"]],
        ["Atendidos 100%", resultado["atendidos_100"]],
        ["Solicitações", resultado["solicitacoes"]],
        ["Entregues em dia", resultado["entregues_em_dia"]],
        ["Projetos atendidos (%)", resultado["projetos_atendidos_pct"]],
        ["Entregas em dia (%)", resultado["entregas_em_dia_pct"]],
        ["ENTREGAS NO PRAZO (%)", resultado["resultado_final"]],
        ["Meta (%)", meta],
        ["", ""],
        ["Entregas diretas (Separação <= CM)", resultado["entrega_direta"]],
        ["Exceção — Última solicitação > CM", resultado["excecao_ultima_solicitacao"]],
        ["Sem separação + Tipo II", resultado["sem_separacao_tipo_ii"]],
        ["Pendências brutas", resultado["pendencias_brutas"]],
        ["Pendências zeradas — sem semana", resultado["pendencias_sem_semana"]],
        ["Pendências zeradas — Tipo II", resultado["pendencias_tipo_ii"]],
        ["Pendências efetivas", resultado["pendencias_efetivas"]],
        ["Códigos sem cadastro", resultado["codigos_sem_cadastro"]],
        ["", ""],
        ["Arquivo RelatorioGeral", resultado["relatorio_nome"]],
        ["SHA-256 RelatorioGeral", resultado["relatorio_sha256"]],
        ["Tamanho RelatorioGeral (bytes)", resultado["relatorio_tamanho_bytes"]],
        ["Arquivo CADASTROS", resultado["cadastros_nome"]],
        ["SHA-256 CADASTROS", resultado["cadastros_sha256"]],
        ["Tamanho CADASTROS (bytes)", resultado["cadastros_tamanho_bytes"]],
    ]

    metodologia_rows = [
        ["ETAPA", "REGRA APLICADA"],
        ["Período-base", "DATA CM entre o primeiro dia do mês da análise e a data da análise, inclusive."],
        ["Programados", "Quantidade de projetos distintos no período-base."],
        ["Solicitações", "Quantidade total de linhas no período-base."],
        ["Atendidos 100%", "Projeto sem nenhuma pendência efetiva após os tratamentos."],
        ["Pendência sem semana", "Pendência > 0 sem Semana de Necessidade numérica é considerada 0."],
        ["Pendência Tipo II", "Pendência > 0 com código classificado como Tipo II em CADASTROS é considerada 0."],
        ["Entrega direta", "Data de Separação menor ou igual à Data CM."],
        ["Exceção de prazo", "Se Data de Separação > Data CM e Última Solicitação > Data CM, a linha é considerada entregue no prazo."],
        ["Sem separação Tipo II", "Linha sem Data de Separação cujo código é Tipo II é considerada entregue no prazo."],
        ["Projetos atendidos (%)", "Atendidos 100% / Programados × 100."],
        ["Entregas em dia (%)", "Entregues em dia / Solicitações × 100."],
        ["Resultado final", "Entregas em dia (%) × Projetos atendidos (%) / 100."],
    ]

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(resumo_rows).to_excel(writer, sheet_name="RESUMO", index=False, header=False)
        pd.DataFrame(resultado["projetos"]).to_excel(writer, sheet_name="PROJETOS", index=False)
        pd.DataFrame(resultado["detalhamento"]).to_excel(writer, sheet_name="DETALHAMENTO", index=False)
        pd.DataFrame(resultado["excecoes"]).to_excel(writer, sheet_name="EXCECOES", index=False)
        pd.DataFrame(resultado["cadastro_usado"]).to_excel(writer, sheet_name="CADASTRO_USADO", index=False)
        pd.DataFrame(metodologia_rows).to_excel(writer, sheet_name="METODOLOGIA", index=False, header=False)

        wb = writer.book
        fill_header = PatternFill("solid", fgColor="1F2937")
        fill_primary = PatternFill("solid", fgColor="FFD43D")
        fill_soft = PatternFill("solid", fgColor="FFF4BF")
        font_white = Font(color="FFFFFF", bold=True)
        font_dark_bold = Font(color="111111", bold=True)
        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for ws in wb.worksheets:
            ws.freeze_panes = "A2" if ws.title != "RESUMO" else "A2"
            ws.sheet_view.showGridLines = False
            max_col = ws.max_column
            max_row = ws.max_row
            if ws.title not in ("RESUMO", "METODOLOGIA") and max_row >= 1:
                for cell in ws[1]:
                    cell.fill = fill_header
                    cell.font = font_white
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border
                ws.auto_filter.ref = ws.dimensions
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if ws.title not in ("RESUMO", "METODOLOGIA"):
                        cell.border = border
            for col_idx in range(1, max_col + 1):
                letter = get_column_letter(col_idx)
                values = [str(ws.cell(r, col_idx).value or "") for r in range(1, min(max_row, 200) + 1)]
                width = min(max(max((len(v) for v in values), default=8) + 2, 10), 42)
                ws.column_dimensions[letter].width = width

        ws = wb["RESUMO"]
        ws.merge_cells("A1:B1")
        ws["A1"].fill = fill_primary
        ws["A1"].font = Font(size=14, bold=True, color="111111")
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.column_dimensions["A"].width = 39
        ws.column_dimensions["B"].width = 78
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).font = Font(bold=True)
            ws.cell(r, 1).border = border
            ws.cell(r, 2).border = border
        for r in [8, 9, 10, 11, 12, 13, 14, 15]:
            ws.cell(r, 1).fill = fill_soft
            ws.cell(r, 2).fill = fill_soft
        for r in [12, 13, 14, 15]:
            ws.cell(r, 2).number_format = "0.00%" if False else "0.00"

        ws = wb["METODOLOGIA"]
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 110
        for cell in ws[1]:
            cell.fill = fill_header
            cell.font = font_white
            cell.border = border
        for row in ws.iter_rows():
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for sheet_name in ["PROJETOS", "DETALHAMENTO", "EXCECOES", "CADASTRO_USADO"]:
            wsx = wb[sheet_name]
            wsx.freeze_panes = "A2"
            wsx.auto_filter.ref = wsx.dimensions

    output.seek(0)
    return output.getvalue()


def _salvar_resultado(resultado, meta):
    client = get_client()
    data_analise = resultado["data_analise"]
    audit_payload = {
        k: resultado[k]
        for k in [
            "periodo_inicio", "periodo_fim", "data_analise", "gerado_em", "programados", "atendidos_100",
            "solicitacoes", "entregues_em_dia", "projetos_atendidos_pct", "entregas_em_dia_pct", "resultado_final",
            "entrega_direta", "excecao_ultima_solicitacao", "sem_separacao_tipo_ii", "pendencias_brutas",
            "pendencias_sem_semana", "pendencias_tipo_ii", "pendencias_efetivas", "codigos_sem_cadastro",
            "relatorio_nome", "relatorio_sha256", "relatorio_tamanho_bytes", "cadastros_nome", "cadastros_sha256",
            "cadastros_tamanho_bytes",
        ]
    }
    payload = {
        "competencia": data_analise,
        "categoria": "OPERACIONAL",
        "indicador": INDICADOR,
        "valor": round(float(resultado["resultado_final"]), 2),
        "meta": round(float(meta), 2),
        "unidade": "%",
        "observacao": json.dumps({"origem": "auditoria_entregas_no_prazo", **audit_payload}, ensure_ascii=False, separators=(",", ":")),
        "atualizado_em": datetime.now(timezone.utc).isoformat(),
    }
    existente = (
        client.table("almox_indicadores").select("id").eq("indicador", INDICADOR)
        .eq("competencia", data_analise).limit(1).execute().data or []
    )
    if existente:
        client.table("almox_indicadores").update(payload).eq("id", existente[0]["id"]).execute()
        acao = "atualizado"
    else:
        client.table("almox_indicadores").insert(payload).execute()
        acao = "salvo"

    try:
        client.table("almox_historico").insert({
            "tipo": "indicador_entregas_no_prazo_auditoria",
            "descricao": f"Entregas no prazo {acao}: {pd.to_datetime(data_analise).strftime('%d/%m/%Y')}",
            "dados": {
                "competencia": data_analise,
                "valor": payload["valor"],
                "meta": payload["meta"],
                "programados": resultado["programados"],
                "atendidos_100": resultado["atendidos_100"],
                "solicitacoes": resultado["solicitacoes"],
                "entregues_em_dia": resultado["entregues_em_dia"],
                "excecao_ultima_solicitacao": resultado["excecao_ultima_solicitacao"],
                "relatorio_sha256": resultado["relatorio_sha256"],
                "cadastros_sha256": resultado["cadastros_sha256"],
            },
        }).execute()
    except Exception:
        pass
    return acao


def render_alimentacao_entregas_auditoria(indicadores):
    hoje = _agora_local().date()
    with st.expander("ALIMENTAR · ENTREGAS NO PRAZO", expanded=False):
        st.caption(
            f"Data da análise fixada automaticamente em {hoje.strftime('%d/%m/%Y')}. "
            "Ao enviar as duas bases, o cálculo é executado automaticamente para o período do dia 01 até hoje."
        )
        a, b, c = st.columns([1.25, 1.25, 0.8])
        with a:
            relatorio = st.file_uploader(
                "RelatorioGeral_Tratado", type=["xlsx", "xlsm", "xltx"], key="audit_entregas_relatorio"
            )
        with b:
            cadastros = st.file_uploader(
                "CADASTROS", type=["xlsx", "xlsm", "xltx"], key="audit_entregas_cadastros"
            )
        with c:
            st.text_input("Data da análise", value=hoje.strftime("%d/%m/%Y"), disabled=True, key="audit_entregas_data")
            meta = st.number_input(
                "Meta (%)", min_value=0.0, max_value=100.0, value=float(_ultima_meta(indicadores)), step=0.1,
                key="audit_entregas_meta",
            )

        if relatorio is None or cadastros is None:
            st.info("Envie as duas planilhas para calcular automaticamente o indicador de hoje.")
            return

        fingerprint = hashlib.sha256(
            relatorio.getvalue() + cadastros.getvalue() + hoje.isoformat().encode("utf-8")
        ).hexdigest()
        if st.session_state.get("audit_entregas_fingerprint") != fingerprint:
            try:
                with st.spinner("Calculando e preparando evidências de auditoria..."):
                    resultado = calcular_entregas_no_prazo_auditoria(relatorio, cadastros, hoje)
                st.session_state["audit_entregas_resultado"] = resultado
                st.session_state["audit_entregas_fingerprint"] = fingerprint
                st.session_state.pop("audit_entregas_registrado", None)
            except Exception as exc:
                st.session_state.pop("audit_entregas_resultado", None)
                st.error(f"Não foi possível calcular o indicador: {exc}")
                return

        resultado = st.session_state.get("audit_entregas_resultado")
        if not resultado:
            return

        st.success(
            f"Apuração concluída para {hoje.strftime('%d/%m/%Y')} — período "
            f"01/{hoje.strftime('%m/%Y')} a {hoje.strftime('%d/%m/%Y')}."
        )

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Programados", resultado["programados"])
        c2.metric("Atendidos 100%", resultado["atendidos_100"])
        c3.metric("Solicitações", resultado["solicitacoes"])
        c4.metric("Entregues em dia", resultado["entregues_em_dia"])

        p1, p2, p3 = st.columns(3)
        p1.metric("Projetos atendidos", f"{resultado['projetos_atendidos_pct']:.2f}%")
        p2.metric("Entregas em dia", f"{resultado['entregas_em_dia_pct']:.2f}%")
        p3.metric("ENTREGAS NO PRAZO", f"{resultado['resultado_final']:.2f}%")

        st.caption(
            f"Evidência de entregas: {resultado['entrega_direta']} diretas + "
            f"{resultado['excecao_ultima_solicitacao']} exceções por Última Solicitação > Data CM + "
            f"{resultado['sem_separacao_tipo_ii']} sem separação / Tipo II."
        )
        st.caption(
            f"Evidência de pendências: {resultado['pendencias_brutas']} brutas; "
            f"{resultado['pendencias_sem_semana']} zeradas por ausência de semana; "
            f"{resultado['pendencias_tipo_ii']} zeradas por Tipo II; "
            f"{resultado['pendencias_efetivas']} efetivas."
        )
        if resultado["codigos_sem_cadastro"]:
            st.warning(f"Há {resultado['codigos_sem_cadastro']} linha(s) com código não localizado em CADASTROS.")

        with st.expander("CONFERÊNCIA POR PROJETO"):
            st.dataframe(pd.DataFrame(resultado["projetos"]), use_container_width=True, hide_index=True)
        with st.expander("DETALHAMENTO LINHA A LINHA"):
            st.dataframe(pd.DataFrame(resultado["detalhamento"]), use_container_width=True, hide_index=True)

        excel_bytes = _gerar_excel_auditoria(resultado, meta)
        b1, b2 = st.columns(2)
        with b1:
            st.download_button(
                "EXPORTAR RELATÓRIO DE AUDITORIA · EXCEL",
                data=excel_bytes,
                file_name=f"auditoria_entregas_no_prazo_{hoje.strftime('%Y-%m-%d')}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                type="secondary",
                key="audit_entregas_excel",
            )
        with b2:
            if st.button(
                f"REGISTRAR RESULTADO DE HOJE · {hoje.strftime('%d/%m/%Y')}",
                type="primary", use_container_width=True, key="audit_entregas_salvar"
            ):
                try:
                    acao = _salvar_resultado(resultado, meta)
                    st.session_state["audit_entregas_registrado"] = True
                    st.success(f"Resultado {acao} no Supabase com data {hoje.strftime('%d/%m/%Y')}: {resultado['resultado_final']:.2f}%.")
                except Exception as exc:
                    st.error(f"Não foi possível registrar o resultado: {exc}")

        if st.session_state.get("audit_entregas_registrado"):
            st.info("O resultado de hoje já foi registrado nesta sessão. Um novo registro na mesma data atualiza o existente, sem duplicar o indicador.")
