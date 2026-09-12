from __future__ import annotations

import json
import math
from datetime import date, datetime
from io import BytesIO

import pandas as pd
import streamlit as st

from supabase_client import get_client

INDICADOR_OTIF = "ENTREGAS NO PRAZO"
SNAPSHOT_TABLE = "almox_indicador_snapshots"


def _json_safe(value):
    if value is None:
        return None
    if isinstance(value, dict):
        return {str(k): _json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_json_safe(v) for v in value]
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if hasattr(value, "item"):
        try:
            return _json_safe(value.item())
        except Exception:
            pass
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _fmt_pct(value):
    try:
        return f"{float(value):.2f}%".replace(".", ",")
    except Exception:
        return "—"


def _fmt_num(value):
    try:
        return f"{float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except Exception:
        return "—"


def _fmt_date(value):
    d = pd.to_datetime(value, errors="coerce")
    return "—" if pd.isna(d) else d.strftime("%d/%m/%Y")


def _fmt_datetime(value):
    d = pd.to_datetime(value, errors="coerce")
    return "—" if pd.isna(d) else d.strftime("%d/%m/%Y %H:%M")


@st.cache_data(ttl=120, show_spinner=False)
def listar_snapshots_otif():
    client = get_client()
    return (
        client.table(SNAPSHOT_TABLE)
        .select("id,indicador,competencia,versao,valor,meta,periodo_inicio,periodo_fim,logic_version,resumo,criado_em")
        .eq("indicador", INDICADOR_OTIF)
        .order("criado_em", desc=True)
        .limit(250)
        .execute()
        .data
        or []
    )


@st.cache_data(ttl=120, show_spinner=False, max_entries=20)
def carregar_snapshot(snapshot_id: str):
    client = get_client()
    rows = (
        client.table(SNAPSHOT_TABLE)
        .select("*")
        .eq("id", snapshot_id)
        .limit(1)
        .execute()
        .data
        or []
    )
    return rows[0] if rows else None


def salvar_snapshot_otif(resultado: dict, meta: float) -> int:
    client = get_client()
    competencia = str(resultado.get("data_registro") or "")
    anteriores = (
        client.table(SNAPSHOT_TABLE)
        .select("versao")
        .eq("indicador", INDICADOR_OTIF)
        .eq("competencia", competencia)
        .order("versao", desc=True)
        .limit(1)
        .execute()
        .data
        or []
    )
    versao = int(anteriores[0]["versao"]) + 1 if anteriores else 1

    resumo = {
        "ontime_almox_pct": resultado.get("ontime_almox_pct"),
        "infull_almox_pct": resultado.get("infull_almox_pct"),
        "otif_almox_pct": resultado.get("otif_almox_pct"),
        "ontime_global_pct": resultado.get("ontime_global_pct"),
        "infull_global_pct": resultado.get("infull_global_pct"),
        "otif_global_pct": resultado.get("otif_global_pct"),
        "ontime_almox_sim_pct": resultado.get("ontime_almox_sim_pct"),
        "ontime_global_sim_pct": resultado.get("ontime_global_sim_pct"),
        "almox_base_ontime_qtd": resultado.get("almox_base_ontime_qtd"),
        "almox_ontime_qtd": resultado.get("almox_ontime_qtd"),
        "projetos_programados": resultado.get("projetos_programados"),
        "projetos_completos_almox": resultado.get("projetos_completos_almox"),
        "projetos_infull_almox": resultado.get("projetos_infull_almox"),
        "linhas_macro": resultado.get("linhas_macro"),
        "solicitacoes_consolidadas": resultado.get("solicitacoes_consolidadas"),
    }

    payload = {
        "indicador": INDICADOR_OTIF,
        "categoria": "OPERACIONAL",
        "competencia": competencia,
        "versao": versao,
        "valor": round(float(resultado.get("otif_almox_pct") or 0), 2),
        "meta": round(float(meta), 2),
        "unidade": "%",
        "periodo_inicio": resultado.get("periodo_inicio"),
        "periodo_fim": resultado.get("periodo_fim"),
        "logic_version": resultado.get("logic_version"),
        "resumo": _json_safe(resumo),
        "detalhes": _json_safe(resultado),
    }
    client.table(SNAPSHOT_TABLE).insert(payload).execute()
    listar_snapshots_otif.clear()
    carregar_snapshot.clear()
    return versao


def _to_frame(rows):
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows)


def _snapshot_label(row):
    return (
        f"{_fmt_date(row.get('competencia'))} · v{row.get('versao')} · "
        f"OTIF {_fmt_pct(row.get('valor'))} · registrado {_fmt_datetime(row.get('criado_em'))}"
    )


def _resumo_rows(resultado, meta, versao):
    return [
        ["AUDITORIA HISTÓRICA · OTIF ALMOX", ""],
        ["Versão do snapshot", versao],
        ["Data de registro", _fmt_date(resultado.get("data_registro"))],
        ["Período", f"{_fmt_date(resultado.get('periodo_inicio'))} a {_fmt_date(resultado.get('periodo_fim'))}"],
        ["Versão da lógica", resultado.get("logic_version") or ""],
        ["ON TIME ALMOX (%)", resultado.get("ontime_almox_pct")],
        ["IN FULL ALMOX (%)", resultado.get("infull_almox_pct")],
        ["OTIF ALMOX (%)", resultado.get("otif_almox_pct")],
        ["Meta (%)", meta],
        ["ON TIME GLOBAL (%)", resultado.get("ontime_global_pct")],
        ["IN FULL GLOBAL (%)", resultado.get("infull_global_pct")],
        ["OTIF GLOBAL (%)", resultado.get("otif_global_pct")],
        ["Base elegível On Time Almox", resultado.get("almox_base_ontime_qtd")],
        ["Atendida no prazo Almox", resultado.get("almox_ontime_qtd")],
        ["Projetos programados", resultado.get("projetos_programados")],
        ["Solicitações consolidadas", resultado.get("solicitacoes_consolidadas")],
    ]


@st.cache_data(show_spinner=False, max_entries=8)
def _excel_snapshot(payload_json: str, meta: float, versao: int):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    resultado = json.loads(payload_json)
    output = BytesIO()
    abas = {
        "CAUSAS": resultado.get("causas", []),
        "PROJETOS": resultado.get("projetos", []),
        "SOLICITACOES": resultado.get("solicitacoes", []),
        "FUNIL_GERAL": resultado.get("funil_geral", []),
        "MATERIAIS_MRP": resultado.get("materiais_mrp", []),
        "DATAS_CM": resultado.get("datas_projeto", []),
    }
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(_resumo_rows(resultado, meta, versao)).to_excel(writer, sheet_name="RESUMO", index=False, header=False)
        for nome, rows in abas.items():
            _to_frame(rows).to_excel(writer, sheet_name=nome, index=False)

        wb = writer.book
        dark = PatternFill("solid", fgColor="1F2937")
        yellow = PatternFill("solid", fgColor="FFD43D")
        white = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D1D5DB")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for ws in wb.worksheets:
            ws.sheet_view.showGridLines = False
            if ws.title != "RESUMO" and ws.max_row >= 1:
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.fill = dark
                    cell.font = white
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    cell.border = border
            for col in range(1, ws.max_column + 1):
                vals = [str(ws.cell(r, col).value or "") for r in range(1, min(ws.max_row, 120) + 1)]
                ws.column_dimensions[get_column_letter(col)].width = min(max(max((len(v) for v in vals), default=8) + 2, 10), 42)
        ws = wb["RESUMO"]
        ws.merge_cells("A1:B1")
        ws["A1"].fill = yellow
        ws["A1"].font = Font(size=14, bold=True, color="111111")
        ws.column_dimensions["A"].width = 42
        ws.column_dimensions["B"].width = 70
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).font = Font(bold=True)
    output.seek(0)
    return output.getvalue()


@st.fragment
def render_historico_otif():
    with st.expander("CONSULTAR · HISTÓRICO DETALHADO OTIF", expanded=False):
        st.caption(
            "Cada registro confirmado gera uma versão imutável da memória de cálculo. "
            "A consulta abaixo não depende de reenviar os relatórios de origem."
        )
        try:
            snapshots = listar_snapshots_otif()
        except Exception as exc:
            st.error(f"Não foi possível consultar o histórico detalhado: {exc}")
            return
        if not snapshots:
            st.info("Ainda não existem snapshots detalhados. O próximo registro do OTIF criará a versão 1.")
            return

        opcoes = {str(r["id"]): _snapshot_label(r) for r in snapshots}
        snapshot_id = st.selectbox(
            "Registro histórico",
            options=list(opcoes.keys()),
            format_func=lambda x: opcoes[x],
            key="otif_snapshot_id",
        )
        try:
            snap = carregar_snapshot(snapshot_id)
        except Exception as exc:
            st.error(f"Não foi possível carregar o snapshot: {exc}")
            return
        if not snap:
            st.warning("Snapshot não encontrado.")
            return

        resultado = snap.get("detalhes") or {}
        meta = float(snap.get("meta") or 0)
        st.caption(
            f"Período: {_fmt_date(snap.get('periodo_inicio'))} a {_fmt_date(snap.get('periodo_fim'))} · "
            f"Snapshot v{snap.get('versao')} · lógica {snap.get('logic_version') or '—'}"
        )

        a, b, c, d = st.columns(4)
        a.metric("ON TIME ALMOX", _fmt_pct(resultado.get("ontime_almox_pct")))
        b.metric("IN FULL ALMOX", _fmt_pct(resultado.get("infull_almox_pct")))
        c.metric("OTIF ALMOX", _fmt_pct(resultado.get("otif_almox_pct")))
        d.metric("META", _fmt_pct(meta))

        g1, g2, g3, g4 = st.columns(4)
        g1.metric("ON TIME GLOBAL", _fmt_pct(resultado.get("ontime_global_pct")))
        g2.metric("IN FULL GLOBAL", _fmt_pct(resultado.get("infull_global_pct")))
        g3.metric("OTIF GLOBAL", _fmt_pct(resultado.get("otif_global_pct")))
        g4.metric("BASE ELEGÍVEL", _fmt_num(resultado.get("almox_base_ontime_qtd")))

        secoes = {
            "Resumo de causas": "causas",
            "Projetos / In Full": "projetos",
            "Solicitações / On Time": "solicitacoes",
            "Funil geral": "funil_geral",
            "Materiais x MRP": "materiais_mrp",
            "Datas CM": "datas_projeto",
        }
        secao = st.radio(
            "Detalhamento",
            options=list(secoes.keys()),
            horizontal=True,
            key="otif_snapshot_secao",
        )
        df = _to_frame(resultado.get(secoes[secao], []))
        if df.empty:
            st.info("Esta seção não possui registros neste snapshot.")
        else:
            st.dataframe(df, use_container_width=True, hide_index=True, height=430)

        payload_json = json.dumps(resultado, ensure_ascii=False, sort_keys=True, default=str)
        excel = _excel_snapshot(payload_json, meta, int(snap.get("versao") or 1))
        st.download_button(
            "EXPORTAR ESTE SNAPSHOT · EXCEL",
            data=excel,
            file_name=f"historico_otif_{snap.get('competencia')}_v{snap.get('versao')}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            key=f"otif_snapshot_excel_{snapshot_id}",
        )
